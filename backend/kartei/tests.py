import base64
import json
import os
import re
import tempfile
import uuid
from datetime import timedelta
from unittest import mock

from django.contrib.auth.hashers import make_password
from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.utils import timezone

from pathlib import Path

from . import batch, exports, graph
from .api import import_fields
from .importer import import_csv
from .models import AuditEntry, ExportProfile, FieldDef, Profile, Proposal, Record, Review, Role, Settings, ref_names
from .vcard import fold, render

PERSON_UID = uuid.UUID("ad081c47-96f6-4b88-9d9b-67faf688772d")
COMPANY_UID = uuid.UUID("cb114125-b78f-43fd-85a0-4fcdea9a8cac")

# Reference vCards for iOS and Android (PHOTO/LOGO/SOURCE excluded: data-dependent)
ANDROID_PERSON = r"""UID:urn:uuid:ad081c47-96f6-4b88-9d9b-67faf688772d
N:Meier;Alexander;Ralph;Prof. Dr. Dr.;
FN:Prof. Dr. Dr. Alexander Meier
X-PHONETIC-FIRST-NAME:Ahlecksanda
X-PHONETIC-MIDDLE-NAME:Ralf
X-PHONETIC-LAST-NAME:Maiher
NICKNAME:Alex
TITLE:Rechenzentrums- & Systembetrieb
ORG:Muster IT GmbH
ROLE:Informatik
EMAIL;WORK:a.meier@acme.com
EMAIL;HOME:alex.meier@gmx.de
EMAIL;Sonstige:a.meier@gmail.com
TEL;CELL:01701234567
TEL;HOME:050311234567
TEL;WORK:05111234567
TEL;X-Assistenz:05111234568
TEL;WORK;FAX:051112345671117
TEL;HOME;FAX:0267612312311
URL:https://acme.com
BDAY:1994-01-01
ADR;WORK:;Empfang;Muster Straße 1;Hannover;;30159;Deutschland
ADR;HOME:;3. OG;Muster Straße 24a;Garbsen;;30823;Deutschland
X-ANDROID-CUSTOM:vnd.android.cursor.item/relation;Friedrich März;1;;;;;;;;;;;;;
X-ANDROID-CUSTOM:vnd.android.cursor.item/relation;Donald Trumpet;7;;;;;;;;;;;;;
NOTE:Trink seinen Kaffe mit Milch\nAbteilung: Informatik\n\nAnsprechpartner intern: KK\ninterne Kontakte: KA\,ZZ"""

IOS_PERSON = r"""UID:urn:uuid:ad081c47-96f6-4b88-9d9b-67faf688772d
N:Meier;Alexander;Ralph;Prof. Dr. Dr.;
FN:Prof. Dr. Dr. Alexander Meier
NICKNAME:Alex
X-PHONETIC-FIRST-NAME:Ahlecksanda
X-PRONUNCIATION-FIRST-NAME:Ahlecksanda
X-PHONETIC-MIDDLE-NAME:Ralf
X-PHONETIC-LAST-NAME:Maiher
X-PRONUNCIATION-LAST-NAME:Maiher
TITLE:Rechenzentrums- & Systembetrieb
ORG:Muster IT GmbH;Informatik
EMAIL;type=WORK;type=pref:a.meier@acme.com
EMAIL;type=HOME:alex.meier@gmx.de
EMAIL;type=Sonstige:a.meier@gmail.com
TEL;type=CELL;type=VOICE:01701234567
TEL;type=HOME;type=VOICE:050311234567
TEL;type=WORK;type=VOICE;type=pref:05111234567
item1.TEL:05111234568
item1.X-ABLabel:_$!<AssistantPhone>!$_
TEL;type=WORK;type=FAX:051112345671117
TEL;type=HOME;type=FAX:0267612312311
item2.URL;type=pref:https://acme.com
item2.X-ABLabel:_$!<HomePage>!$_
BDAY:1994-01-01
ADR;WORK:;Empfang;Muster Straße 1;Hannover;;30159;Deutschland
ADR;HOME:;3. OG;Muster Straße 24a;Garbsen;;30823;Deutschland
item3.X-ABRELATEDNAMES;type=pref:Friedrich März
item3.X-ABLabel:_$!<Assistant>!$_
item4.X-ABRELATEDNAMES;type=pref:Donald Trumpet
item4.X-ABLabel:_$!<Manager>!$_
NOTE:Trink seinen Kaffe mit Milch\n\nAnsprechpartner intern: KK\ninterne Kontakte: KA\,ZZ"""

# Companies use the same note layout as persons.
ANDROID_COMPANY = r"""ORG:ACME GmbH
NICKNAME:ACME
N:ACME GmbH;;;;
FN:ACME GmbH
EMAIL;Zentrale:info@acme.com
EMAIL;Datenschutz:datenschutz@acme.com
EMAIL;Rechnung:einkauf@acme.com
TEL;Zentrale:05111231232
TEL;Zentrale;FAX:0511123123123
URL:https://acme.com
ADR;WORK:;;Muster Straße 1;Hannover;;30159;Deutschland
X-ANDROID-CUSTOM:vnd.android.cursor.item/relation;Angela Merklin;0;Geschäftsführer;;;;;;;;;;;;
UID:urn:uuid:cb114125-b78f-43fd-85a0-4fcdea9a8cac
NOTE:Test\n\nAnsprechpartner intern: KK\ninterne Kontakte: KA\,UI"""

IOS_COMPANY = r"""ORG:ACME GmbH
NICKNAME:ACME
N:;;;;
FN:ACME GmbH
EMAIL;type=Zentrale;type=pref:info@acme.com
EMAIL;type=Datenschutz:datenschutz@acme.com
EMAIL;type=Rechnung:einkauf@acme.com
TEL;type=Zentrale;type=VOICE;type=pref:05111231232
TEL;type=Zentrale;type=FAX:0511123123123
item1.URL;type=pref:https://acme.com
item1.X-ABLabel:_$!<HomePage>!$_
ADR;WORK:;;Muster Straße 1;Hannover;;30159;Deutschland
item2.X-ABRELATEDNAMES;type=pref:Angela Merklin
item2.X-ABLabel:Geschäftsführer
UID:urn:uuid:cb114125-b78f-43fd-85a0-4fcdea9a8cac
NOTE:Test\n\nAnsprechpartner intern: KK\ninterne Kontakte: KA\,UI"""


EXAMPLES = Path(__file__).resolve().parents[2] / "examples"
# A CSV export layout: (header, person source, company source)
EXPORT_COLUMNS = [
    ("ID", "calc:legacy_or_uid", "calc:legacy_or_uid"),
    ("Anrede", "field:salutation", ""),
    ("Vorname", "core:first_name", "field:short_name"),
    ("Nachname", "core:last_name", ""),
    ("Funktion", "field:job_title", ""),
    ("Abteilung", "field:department", ""),
    ("Firma", "calc:company_name", "core:name"),
    ("Straße", "field:work_street", "field:work_street"),
    ("PLZ", "field:work_zip", "field:work_zip"),
    ("Ort", "field:work_city", "field:work_city"),
    ("Land", "field:work_country", "field:work_country"),
    ("Telefon", "field:phone_work", "field:company_phone"),
    ("Fax", "field:fax", "field:company_fax"),
    ("Mobil", "field:mobile", ""),
    ("E-Mail", "field:email", "field:company_email"),
    ("Ansprechpartner intern", "field:internal_contact", "field:internal_contact"),
    ("Kontakte intern", "field:internal_contacts", "field:internal_contacts"),
    ("Schlagworte", "field:keywords", "field:keywords"),
    ("Notizen", "field:description", "field:description"),
]
# Import mapping for the test CSVs: columns map by suggestion; companies put their short name into "Vorname",
# IDs "F-…" are companies, "P-…" persons.
IMPORT_MAPPING = {"format": "kartei-import-mapping", "version": 1,
                  "columns": {"Vorname": {"person": "core:first_name", "company": "field:short_name"}},
                  "kind": {"column": "", "company_prefix": "F-", "person_prefix": "P-"}}


def example(name):
    return json.loads((EXAMPLES / name).read_text(encoding="utf-8"))


def load_fields(name="fields-example.json"):
    """A full field set through the field import."""
    res = import_fields(example(name), apply=True)["fields"]
    assert all(r["status"] == "new" for r in res), res


def import_test_csv(raw, dry_run=True):
    return import_csv(raw, mapping=IMPORT_MAPPING, dry_run=dry_run)


def card_lines(vcf):
    lines = vcf.replace("\r\n ", "").split("\r\n")
    return {l for l in lines if l and not l.startswith(("BEGIN:", "VERSION:", "END:", "SOURCE:", "PHOTO", "LOGO"))}


def sample_records():
    load_fields()
    muster = Record.objects.create(kind="company", name="Muster IT GmbH")
    person = Record.objects.create(
        uid=PERSON_UID, kind="person", first_name="Alexander", last_name="Meier", company=muster, data={
            "academic_title": "Prof. Dr. Dr.", "middle_name": "Ralph", "nickname": "Alex",
            "phonetic_first": "Ahlecksanda", "phonetic_middle": "Ralf", "phonetic_last": "Maiher",
            "job_title": "Rechenzentrums- & Systembetrieb", "department": "Informatik",
            "email": "a.meier@acme.com", "email_private": "alex.meier@gmx.de", "email_other": "a.meier@gmail.com",
            "mobile": "01701234567", "phone_home": "050311234567", "phone_work": "05111234567",
            "assistant_phone": "05111234568", "fax": "051112345671117", "fax_home": "0267612312311",
            "website": "https://acme.com", "birthday": "1994-01-01",
            "work_street": "Muster Straße 1", "work_extra": "Empfang", "work_city": "Hannover", "work_zip": "30159",
            "work_country": "Deutschland", "home_street": "Muster Straße 24a", "home_extra": "3. OG",
            "home_city": "Garbsen", "home_zip": "30823", "home_country": "Deutschland",
            "assistant": "Friedrich März", "manager": "Donald Trumpet", "internal_contact": "KK",
            "internal_contacts": "KA,ZZ", "description": "Trink seinen Kaffe mit Milch"})
    acme = Record.objects.create(
        uid=COMPANY_UID, kind="company", name="ACME GmbH", data={
            "short_name": "ACME", "company_email": "info@acme.com", "email_privacy": "datenschutz@acme.com",
            "email_billing": "einkauf@acme.com", "company_phone": "05111231232", "company_fax": "0511123123123",
            "website": "https://acme.com", "work_street": "Muster Straße 1", "work_city": "Hannover",
            "work_zip": "30159", "work_country": "Deutschland", "ceo": "Angela Merklin",
            "internal_contact": "KK", "internal_contacts": "KA,UI", "description": "Test"})
    return person, acme


class VCardTests(TestCase):
    def setUp(self):
        self.person, self.company = sample_records()
        self.fields = list(FieldDef.objects.all())

    def test_reference_cards(self):
        for rec, flavour, expected in [(self.person, "android", ANDROID_PERSON), (self.person, "ios", IOS_PERSON),
                                       (self.company, "android", ANDROID_COMPANY), (self.company, "ios", IOS_COMPANY)]:
            with self.subTest(rec=rec.display_name, flavour=flavour):
                got = card_lines(render(rec, self.fields, flavour, "https://kartei.example"))
                self.assertEqual(got, set(expected.splitlines()))

    def test_fold_and_escape(self):
        line = "NOTE:" + "ä" * 100
        folded = fold(line)
        self.assertTrue(all(len(l.encode()) <= 75 for l in folded.split("\r\n")))
        self.assertEqual(folded.replace("\r\n ", ""), line)
        self.person.data["description"] = "a;b,c\\d\nzweite Zeile"
        vcf = render(self.person, self.fields, "ios")
        self.assertIn(r"NOTE:a\;b\,c\\d\nzweite Zeile", vcf.replace("\r\n ", ""))


class ApiTestBase(TestCase):
    def setUp(self):
        Settings.get()
        load_fields()
        FieldDef.objects.filter(key="internal_contact").update(required=False)  # ReviewTests cover the required case
        self.users = {}
        for role in Role.values:
            u = User.objects.create_user(role, password="pw-" + role * 3)
            Profile.objects.create(user=u, role=role)
            self.users[role] = u

    def as_(self, role):
        self.client.force_login(self.users[role])

    def j(self, method, url, body=None, status=200):
        resp = getattr(self.client, method)(url, json.dumps(body) if body is not None else "",
                                            content_type="application/json")
        self.assertEqual(resp.status_code, status, resp.content[:500])
        return resp.json() if resp.get("Content-Type", "").startswith("application/json") else resp


class RecordApiTests(ApiTestBase):
    def test_roles(self):
        self.client.logout()
        self.j("get", "/api/bootstrap", status=401)
        self.as_("read")
        self.j("post", "/api/records", {"kind": "company", "name": "X"}, status=403)
        self.j("get", "/api/admin/fields", status=403)
        self.as_("maintainer")
        self.j("post", "/api/records", {"kind": "company", "name": "X"}, status=201)
        self.j("get", "/api/admin/fields", status=403)
        self.as_("admin")
        self.j("get", "/api/admin/fields")

    def test_create_update_conflict_audit(self):
        self.as_("maintainer")
        c = self.j("post", "/api/records", {"kind": "company", "name": "ACME", "data": {"company_phone": "0511 1"}},
                   status=201)
        p = self.j("post", "/api/records", {"kind": "person", "last_name": "Meier", "company": c["uid"],
                                            "data": {"email": "a@b.de", "keywords": []}}, status=201)
        self.assertEqual(p["company"], c["uid"])
        p2 = self.j("patch", f"/api/records/{p['uid']}", {"version": p["version"], "data": {"mobile": "+49 170 1"}})
        self.assertEqual(p2["version"], p["version"] + 1)
        stale = self.j("patch", f"/api/records/{p['uid']}", {"version": p["version"], "data": {"mobile": "1"}},
                       status=409)
        self.assertEqual(stale["record"]["data"]["mobile"], "+49 170 1")
        entry = AuditEntry.objects.filter(target_id=p["uid"], action="update").first()
        self.assertEqual(entry.changes, {"mobile": [None, "+49 170 1"]})
        self.assertEqual(entry.username, "maintainer")

    def test_validation(self):
        self.as_("maintainer")
        FieldDef.objects.create(key="count", label_de="Anzahl", label_en="Count", applies_to="both", type="integer",
                                min_value=0, max_value=10)
        FieldDef.objects.create(key="rate", label_de="Satz", label_en="Rate", applies_to="both", type="decimal",
                                decimals=2)
        FieldDef.objects.create(key="tags", label_de="Tags", label_en="Tags", applies_to="both", type="multiselect",
                                options=[{"value": "a", "de": "A", "en": "A"}, {"value": "b", "de": "B", "en": "B"}])
        bad = self.j("post", "/api/records", {"kind": "person", "last_name": "X", "data": {
            "count": "11", "rate": "abc", "tags": ["c"], "email": "nope", "mobile": "abc", "nope": 1}}, status=400)
        self.assertEqual(bad["fields"], {"count": "too_large", "rate": "invalid_number", "tags": "invalid_option",
                                         "email": "invalid_email", "mobile": "invalid_phone", "nope": "unknown_field"})
        ok = self.j("post", "/api/records", {"kind": "person", "last_name": "X", "data": {
            "count": "3", "rate": "1.234,5", "tags": ["b", "a", "b"], "work_zip": "AB-1234", "website": "acme.com"}},
            status=201)
        self.assertEqual(ok["data"], {"count": 3, "rate": "1234.50", "tags": ["b", "a"], "work_zip": "AB-1234",
                                      "website": "https://acme.com"})
        self.j("post", "/api/records", {"kind": "person", "data": {}}, status=400)

    def test_field_access(self):
        f = FieldDef.objects.get(key="contract_no")
        f.visible_to, f.editable_by = "maintainer", "admin"
        f.save()
        r = Record.objects.create(kind="person", last_name="Y", data={"contract_no": "secret"})
        self.as_("read")
        rec = next(x for x in self.j("get", "/api/bootstrap")["records"] if x["uid"] == str(r.uid))
        self.assertNotIn("contract_no", rec["data"])
        self.as_("maintainer")
        rec = self.j("get", f"/api/records/{r.uid}")
        self.assertEqual(rec["data"]["contract_no"], "secret")
        bad = self.j("patch", f"/api/records/{r.uid}", {"version": 1, "data": {"contract_no": "x"}}, status=400)
        self.assertEqual(bad["fields"], {"contract_no": "forbidden"})
        self.j("patch", f"/api/records/{r.uid}", {"version": 1, "data": {"contract_no": "secret", "mobile": "1 2 3"}})

    def test_recycle_bin(self):
        self.as_("maintainer")
        r = Record.objects.create(kind="company", name="Gone")
        boot = self.j("get", "/api/bootstrap")
        self.j("delete", f"/api/records/{r.uid}")
        self.assertNotIn(str(r.uid), [x["uid"] for x in self.j("get", "/api/bootstrap")["records"]])
        ch = self.j("get", f"/api/changes?since={boot['cursor'].replace('+', '%2B')}")
        self.assertIn(str(r.uid), ch["deleted"])
        self.assertEqual([x["uid"] for x in self.j("get", "/api/trash")["records"]], [str(r.uid)])
        self.j("delete", f"/api/trash/{r.uid}", status=403)  # purge is admin-only
        self.j("post", f"/api/records/{r.uid}/restore")
        self.assertIn(str(r.uid), [x["uid"] for x in self.j("get", "/api/bootstrap")["records"]])
        actions = list(AuditEntry.objects.filter(target_id=str(r.uid)).values_list("action", flat=True))
        self.assertEqual(actions, ["restore", "delete"])
        self.j("delete", f"/api/records/{r.uid}")
        self.as_("admin")
        self.j("delete", f"/api/trash/{r.uid}")
        self.assertFalse(Record.objects.filter(uid=r.uid).exists())
        self.assertTrue(AuditEntry.objects.filter(target_id=str(r.uid), action="purge").exists())

    def test_image(self):
        self.as_("maintainer")
        r = Record.objects.create(kind="company", name="Img")
        resp = self.client.put(f"/api/img/{r.uid}", b"not an image", content_type="image/png")
        self.assertEqual(resp.status_code, 400)
        png = b"\x89PNG\r\n\x1a\n" + b"0" * 20
        resp = self.client.put(f"/api/img/{r.uid}", png, content_type="image/png")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(self.client.get(f"/api/img/{r.uid}").content, png)


class AdminApiTests(ApiTestBase):
    def test_fields_crud(self):
        self.as_("admin")
        f = self.j("post", "/api/admin/fields", {"label": {"de": "Größe", "en": "Size"}, "applies_to": "person",
                                                 "type": "select", "options": [{"value": "S"}, {"value": "M"}]},
                   status=201)
        self.assertEqual(f["key"], "size")
        self.j("patch", f"/api/admin/fields/{f['id']}", {"type": "multiselect"})
        Record.objects.create(kind="person", last_name="Z", data={"size": ["S"]})
        bad = self.j("patch", f"/api/admin/fields/{f['id']}", {"type": "date"}, status=400)
        self.assertEqual(bad["fields"]["type"], "incompatible_type")
        self.j("delete", f"/api/admin/fields/{f['id']}", status=409)
        self.j("patch", f"/api/admin/fields/{f['id']}", {"archived": True})
        self.assertTrue(AuditEntry.objects.filter(target_type="field", action="update").exists())

    def test_field_regex_and_help(self):
        self.as_("admin")
        bad = self.j("post", "/api/admin/fields", {"label": {"de": "Kd-Nr", "en": "Cust no"}, "applies_to": "both",
                                                   "type": "text", "regex": "([A-Z"}, status=400)
        self.assertEqual(bad["fields"], {"regex": "invalid_regex"})
        f = self.j("post", "/api/admin/fields", {
            "label": {"de": "Kd-Nr", "en": "Cust no"}, "applies_to": "both", "type": "text", "regex": r"K-\d{5}",
            "help": {"de": "Format K-12345", "en": "Format K-12345"}}, status=201)
        self.assertEqual((f["regex"], f["help"]["de"]), (r"K-\d{5}", "Format K-12345"))
        n = self.j("post", "/api/admin/fields", {"label": {"de": "Zahl", "en": "Num"}, "applies_to": "both",
                                                 "type": "integer", "regex": "x"}, status=201)
        self.assertEqual(n["regex"], "")  # regex only applies to text-like types
        self.as_("maintainer")
        err = self.j("post", "/api/records", {"kind": "company", "name": "A", "data": {f["key"]: "K-123"}}, status=400)
        self.assertEqual(err["fields"], {f["key"]: "pattern_mismatch"})
        self.j("post", "/api/records", {"kind": "company", "name": "A", "data": {f["key"]: "K-12345"}}, status=201)

    def test_users_and_last_admin(self):
        self.as_("admin")
        u = self.j("post", "/api/admin/users", {"username": "neu", "password": "long-enough-pw", "role": "read"},
                   status=201)
        self.j("patch", f"/api/admin/users/{u['id']}", {"role": "maintainer"})
        self.assertEqual(Profile.objects.get(user_id=u["id"]).role, "maintainer")
        self.j("patch", f"/api/admin/users/{self.users['admin'].id}", {"role": "read"}, status=409)
        self.j("delete", f"/api/admin/users/{self.users['admin'].id}", status=409)

    def test_settings_secrets(self):
        self.as_("admin")
        s = self.j("put", "/api/admin/settings", {"ldap_uri": "ldaps://dc", "ldap_bind_password": "geheim"})
        self.assertTrue(s["ldap_bind_password_set"])
        self.assertNotIn("geheim", json.dumps(s))
        self.assertEqual(Settings.get().ldap_bind_password, "geheim")
        self.assertNotIn("geheim", json.dumps(list(AuditEntry.objects.values_list("changes", flat=True))))
        r = self.j("put", "/api/admin/settings", {"graph_enabled": True, "graph_tenant_id": "t"}, status=400)
        self.assertEqual(r["fields"], {"graph_enabled": "graph_incomplete"})
        r = self.j("put", "/api/admin/settings", {"graph_folder_responsible": " kartei "}, status=400)
        self.assertEqual(r["fields"], {"graph_folder_responsible": "duplicate"})
        self.j("put", "/api/admin/settings", {"graph_folder_responsible": " Meine ", "graph_folder_internal": "Mine"})
        self.assertEqual(Settings.get().graph_folder_responsible, "Meine")
        r = self.j("put", "/api/admin/settings", {"graph_folder_responsible": "mine "}, status=400)
        self.assertEqual(r["fields"], {"graph_folder_internal": "duplicate"})
        s = self.j("put", "/api/admin/settings", {"graph_enabled": True, "graph_tenant_id": "t", "graph_client_id": "c",
                                                  "graph_group_id": "g", "graph_client_secret": "graphpw"})
        self.assertTrue(s["graph_client_secret_set"])
        self.assertNotIn("graphpw", json.dumps(s) + json.dumps(list(AuditEntry.objects.values_list("changes", flat=True))))
        self.assertEqual(Settings.get().graph_client_secret, "graphpw")
        self.assertTrue(self.j("post", "/api/admin/graph/sync")["graph_sync_requested"])
        self.j("put", "/api/admin/settings", {"graph_enabled": False})
        self.j("post", "/api/admin/graph/sync", status=409)
        self.as_("maintainer")
        self.j("post", "/api/admin/graph/sync", status=403)

    def test_login_and_language(self):
        self.client.logout()
        self.j("post", "/api/login", {"username": "read", "password": "wrong"}, status=401)
        self.j("post", "/api/login", {"username": "read", "password": "pw-" + "read" * 3})
        me = self.j("patch", "/api/me", {"language": "en"})
        self.assertEqual(me["user"]["language"], "en")


class PostalCodeTests(ApiTestBase):
    """PLZ/ZIP must stay text: leading zeros and letters survive every path."""
    CODES = ["01067", "SW1A 1AA", "1010", "A-1010", "D-10115"]

    def test_preset_fields_are_text(self):
        self.assertEqual(set(FieldDef.objects.filter(key__endswith="_zip").values_list("type", flat=True)), {"text"})

    def test_roundtrip(self):
        from openpyxl import load_workbook
        self.as_("maintainer")
        for code in self.CODES:
            r = self.j("post", "/api/records", {"kind": "company", "name": f"Firma {code}", "data": {"work_zip": code}},
                       status=201)
            self.assertEqual(r["data"]["work_zip"], code)
            rec = Record.objects.get(uid=r["uid"])
            self.assertIn(f";{code};", render(rec, list(FieldDef.objects.all()), "ios").split("ADR;WORK:")[1].split("\r\n")[0])
        d = tempfile.mkdtemp()
        with mock.patch.object(exports, "EXPORT_ROOT", d):
            cols = [{"header": "PLZ", "person": "field:work_zip", "company": "field:work_zip"}]
            for fmt in ("csv", "xlsx"):
                p = ExportProfile.objects.create(name=fmt, format=fmt, path=f"{d}/plz.{fmt}", columns=cols)
                self.assertEqual(exports.run(p), "ok", p.last_message)
            csv_text = open(f"{d}/plz.csv", encoding="utf-8-sig").read()
            ws = load_workbook(f"{d}/plz.xlsx").active
        xlsx_values = [c.value for c in ws["A"][1:]]
        for code in self.CODES:
            self.assertIn(code, csv_text.splitlines())
            self.assertIn(code, xlsx_values)  # stored as text cell, not number
        self.assertTrue(all(isinstance(v, str) for v in xlsx_values))

    def test_import_keeps_leading_zero(self):
        csv_data = "ID;Vorname;Nachname;Firma;PLZ\nP-1;A;B;;01067\n".encode()
        import_test_csv(csv_data, dry_run=False)
        self.assertEqual(Record.objects.get(legacy_id="P-1").data["work_zip"], "01067")


class SessionTests(ApiTestBase):
    @override_settings(KARTEI_VERSION="9.9.9-test", KARTEI_SOURCE_URL="https://example.org/fork")
    def test_version_and_source_only_when_signed_in(self):
        self.assertNotIn("version", self.client.get("/api/session").json())
        self.as_("read")
        s = self.client.get("/api/session").json()
        self.assertEqual((s["version"], s["source"]), ("9.9.9-test", "https://example.org/fork"))

    def test_a_deactivated_user_loses_an_open_session(self):
        for backend in ("kartei.auth.DbLDAPBackend", "django.contrib.auth.backends.ModelBackend"):
            u = self.users["maintainer"]
            User.objects.filter(pk=u.pk).update(is_active=True)
            self.client.force_login(u, backend=backend)
            self.j("get", "/api/bootstrap")
            User.objects.filter(pk=u.pk).update(is_active=False)  # e.g. by the LDAP user import
            self.j("get", "/api/bootstrap", status=401)


class UserFieldTests(ApiTestBase):
    """Internal responsible (user) and internal contacts (users) reference Kartei users by username."""
    def setUp(self):
        super().setUp()
        User.objects.filter(username="read").update(first_name="Rita", last_name="Reader")
        User.objects.filter(username="admin").update(first_name="Anna", last_name="Admin")
        self.as_("maintainer")

    def create(self, **data):
        return self.j("post", "/api/records", {"kind": "person", "last_name": "X", "data": data}, status=201)

    def test_options_validation_and_deactivation(self):
        boot = self.j("get", "/api/bootstrap")
        f = next(f for f in boot["fields"] if f["key"] == "internal_contact")
        self.assertEqual(f["type"], "user")
        self.assertIn({"value": "read", "de": "Rita Reader", "en": "Rita Reader"}, f["options"])
        r = self.create(internal_contact="read", internal_contacts=["maintainer", "read", "maintainer"])
        self.assertEqual(r["data"]["internal_contacts"], ["maintainer", "read"])
        # admins are listed (for existing values) but can't be picked
        self.assertIn({"value": "admin", "de": "Anna Admin", "en": "Anna Admin", "pick": False}, f["options"])
        for key, v in (("internal_contact", "admin"), ("internal_contacts", ["admin"])):
            bad = self.j("post", "/api/records", {"kind": "person", "last_name": "A", "data": {key: v}}, status=400)
            self.assertEqual(bad["fields"][key], "invalid_user")
        rec = Record.objects.get(uid=r["uid"])
        rec.data["internal_contacts"] = ["admin", "read"]  # assigned before the rule: stays on unrelated edits
        rec.save()
        r = self.j("patch", f"/api/records/{r['uid']}", {"version": r["version"], "first_name": "Ann"})
        self.assertEqual(r["data"]["internal_contacts"], ["admin", "read"])
        bad = self.j("post", "/api/records", {"kind": "person", "last_name": "Y", "data": {"internal_contact": "nobody"}},
                     status=400)
        self.assertEqual(bad["fields"]["internal_contact"], "invalid_user")
        User.objects.filter(username="read").update(is_active=False)
        url = f"/api/records/{r['uid']}"
        # unchanged values of a deactivated user may stay, new assignments may not
        r = self.j("patch", url, {"version": r["version"], "data": {"internal_contact": "read",
                                                                    "internal_contacts": ["read"]}})
        bad = self.j("post", "/api/records", {"kind": "person", "last_name": "Z", "data": {"internal_contacts": ["read"]}},
                     status=400)
        self.assertEqual(bad["fields"]["internal_contacts"], "invalid_user")
        f = next(f for f in self.j("get", "/api/bootstrap")["fields"] if f["key"] == "internal_contact")
        self.assertIn({"value": "read", "de": "Rita Reader (inaktiv)", "en": "Rita Reader (inactive)", "pick": False},
                      f["options"])

    def test_names_in_vcard_and_export(self):
        r = self.create(internal_contact="read", internal_contacts=["maintainer"])
        rec = Record.objects.get(uid=r["uid"])
        rec.data["internal_contacts"] = ["admin", "maintainer"]  # an admin assigned before admins were excluded
        rec.save()
        vcf = render(rec, list(FieldDef.objects.all()), "ios", names=ref_names()).replace("\r\n ", "")
        self.assertIn("Ansprechpartner intern: Rita Reader", vcf)
        self.assertIn(r"interne Kontakte: Anna Admin\, maintainer", vcf)
        d = tempfile.mkdtemp()
        with mock.patch.object(exports, "EXPORT_ROOT", d):
            cols = [{"header": "U1", "person": "field:internal_contact", "company": ""},
                    {"header": "U2", "person": "field:internal_contacts", "company": ""}]
            p = ExportProfile.objects.create(name="u", format="csv", path=f"{d}/u.csv", columns=cols)
            self.assertEqual(exports.run(p), "ok", p.last_message)
            self.assertIn("Rita Reader;Anna Admin, maintainer", open(f"{d}/u.csv", encoding="utf-8-sig").read())

    def test_type_change_wraps_values(self):
        r = self.create(internal_contact="read")
        self.as_("admin")
        f = FieldDef.objects.get(key="internal_contact")
        self.j("patch", f"/api/admin/fields/{f.id}", {"type": "users"})
        rec = Record.objects.get(uid=r["uid"])
        self.assertEqual((rec.data["internal_contact"], rec.version), (["read"], r["version"] + 1))
        self.j("patch", f"/api/admin/fields/{f.id}", {"type": "text"}, status=400)

    def test_import_matches_users(self):
        csv_data = ("ID;Vorname;Nachname;Firma;Ansprechpartner intern;Kontakte intern\n"
                    "P-1;A;B;;RR;Anna Admin |##| maintainer, Unbekannt\n").encode()  # admins can't be assigned
        res = import_test_csv(csv_data, dry_run=False)
        data = Record.objects.get(legacy_id="P-1").data
        self.assertEqual((data["internal_contact"], data["internal_contacts"]), ("read", ["maintainer"]))
        self.assertTrue(any(w.startswith("@import_unknown_user") and "Unbekannt" in w and "Anna Admin" in w
                            for w in res["warnings"]))


class OpenOptionsTests(ApiTestBase):
    """Keywords: maintainers add new options inline; closed selects stay closed."""
    def test_inline_options(self):
        self.as_("maintainer")
        kw = FieldDef.objects.get(key="keywords")
        self.assertTrue(kw.open_options)
        r = self.j("post", "/api/records", {"kind": "person", "last_name": "A", "data": {"keywords": ["  Neues   Wort "]}},
                   status=201)
        self.assertEqual(r["data"]["keywords"], ["Neues Wort"])
        kw.refresh_from_db()
        self.assertIn({"value": "Neues Wort", "de": "Neues Wort", "en": "Neues Wort"}, kw.options)
        self.assertTrue(AuditEntry.objects.filter(target_type="field", target_id=str(kw.id)).exists())
        # same text in another case maps onto the existing option
        r = self.j("post", "/api/records", {"kind": "person", "last_name": "B", "data": {"keywords": ["neues wort"]}},
                   status=201)
        self.assertEqual(r["data"]["keywords"], ["Neues Wort"])
        kw.refresh_from_db()
        self.assertEqual(sum(o["value"] == "Neues Wort" for o in kw.options), 1)
        boot = self.j("get", "/api/bootstrap")
        self.assertTrue(next(f for f in boot["fields"] if f["key"] == "keywords")["open_options"])

    def test_closed_select_rejects_new(self):
        FieldDef.objects.create(key="prio", label_de="Prio", label_en="Prio", applies_to="both", type="select",
                                options=[{"value": "A", "de": "A", "en": "A"}])
        self.as_("maintainer")
        bad = self.j("post", "/api/records", {"kind": "person", "last_name": "C", "data": {"prio": "B"}}, status=400)
        self.assertEqual(bad["fields"]["prio"], "invalid_option")
        self.assertEqual(len(FieldDef.objects.get(key="prio").options), 1)


class ContactFieldTests(ApiTestBase):
    """Managing director: a person of the company (stored as uid) or free text."""
    def setUp(self):
        super().setUp()
        self.as_("maintainer")
        self.co = self.j("post", "/api/records", {"kind": "company", "name": "Firma"}, status=201)
        self.p = self.j("post", "/api/records", {"kind": "person", "first_name": "Gerda", "last_name": "Chefin",
                                                 "company": self.co["uid"]}, status=201)

    def set_ceo(self, value, status=200):
        co = self.j("get", f"/api/records/{self.co['uid']}")
        return self.j("patch", f"/api/records/{self.co['uid']}", {"version": co["version"], "data": {"ceo": value}},
                      status=status)

    def test_link_free_text_and_rendering(self):
        self.assertEqual(self.set_ceo(self.p["uid"])["data"]["ceo"], self.p["uid"])
        rec = Record.objects.get(uid=self.co["uid"])
        fields = list(FieldDef.objects.all())
        self.assertIn("Gerda Chefin", render(rec, fields, "ios", names=ref_names()))
        Record.objects.filter(uid=self.p["uid"]).update(first_name="Gerda M.")
        self.assertIn("Gerda M. Chefin", render(rec, fields, "android", names=ref_names()))
        d = tempfile.mkdtemp()
        with mock.patch.object(exports, "EXPORT_ROOT", d):
            cols = [{"header": "GF", "person": "", "company": "field:ceo"}]
            p = ExportProfile.objects.create(name="gf", format="csv", path=f"{d}/gf.csv", columns=cols, kinds="company")
            self.assertEqual(exports.run(p), "ok", p.last_message)
            self.assertIn("Gerda M. Chefin", open(f"{d}/gf.csv", encoding="utf-8-sig").read())
        self.assertEqual(self.set_ceo("  Dr.  Hans   Extern ")["data"]["ceo"], "Dr. Hans Extern")
        rec.refresh_from_db()
        self.assertIn("Dr. Hans Extern", render(rec, fields, "ios", names=ref_names()))

    def test_closed_contact_field_rejects_text(self):
        FieldDef.objects.filter(key="ceo").update(open_options=False)
        self.assertEqual(self.set_ceo("Jemand", status=400)["fields"]["ceo"], "invalid_contact")
        self.assertEqual(self.set_ceo(self.co["uid"], status=400)["fields"]["ceo"], "invalid_contact")  # not a person

    def test_text_field_converts_with_free_text(self):
        FieldDef.objects.filter(key="ceo").update(type="text", open_options=False)
        self.set_ceo("Alter Wert")
        self.as_("admin")
        f = FieldDef.objects.get(key="ceo")
        r = self.j("patch", f"/api/admin/fields/{f.id}", {"type": "contact"})
        self.assertTrue(r["open_options"])
        self.as_("maintainer")
        self.set_ceo("Alter Wert")  # still valid


class ReviewTests(ApiTestBase):
    """Reviews go to the responsible user, who confirms or proposes; maintainers decide."""
    def setUp(self):
        super().setUp()
        FieldDef.objects.filter(key="internal_contact").update(required=True)
        User.objects.filter(username="read").update(email="read@example.com", first_name="Rita")
        self.as_("maintainer")
        self.co = self.j("post", "/api/records", {"kind": "company", "name": "Alt GmbH",
                                                  "data": {"internal_contact": "read"}}, status=201)
        self.p = self.j("post", "/api/records", {"kind": "person", "last_name": "Meier", "company": self.co["uid"],
                                                 "data": {"internal_contact": "read", "email": "m@alt.de"}}, status=201)
        self.orphan = Record.objects.create(kind="person", last_name="Ohne")  # imported without responsible

    def test_responsible_required(self):
        bad = self.j("post", "/api/records", {"kind": "person", "last_name": "X"}, status=400)
        self.assertEqual(bad["fields"]["internal_contact"], "required")

    def test_manual_start_mail_and_confirm(self):
        from django.core import mail
        self.as_("admin")
        res = self.j("post", "/api/admin/reviews", {"scope": "all"})
        self.assertEqual(res, {"created": 2, "already_open": 0, "no_responsible": 1})
        self.assertEqual(self.j("post", "/api/admin/reviews", {"scope": "all"})["already_open"], 2)
        stats = self.j("get", "/api/admin/reviews")
        self.assertEqual((stats["open"], stats["no_responsible"]), (2, 1))
        # digest: one mail per user, only once
        Settings.objects.update(smtp_host="smtp.example", smtp_from="kartei@example.com", base_url="https://kartei.example")
        with mock.patch("kartei.reviews.mail_connection", lambda s: mail.get_connection()):
            from kartei import reviews
            self.assertEqual(reviews.notify(Settings.get()), 1)
            self.assertEqual(reviews.notify(Settings.get()), 0)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("Alt GmbH: https://kartei.example/company/", mail.outbox[0].body)
        self.assertEqual(mail.outbox[0].to, ["read@example.com"])
        # the read-only responsible sees and confirms
        self.as_("read")
        boot = self.j("get", "/api/bootstrap")
        self.assertEqual({rv["record"] for rv in boot["reviews"]}, {self.co["uid"], self.p["uid"]})
        since = timezone.now() - timedelta(seconds=1)
        Record.objects.filter(uid=self.co["uid"]).update(updated_at=timezone.now() - timedelta(days=1))
        r = self.j("post", f"/api/records/{self.co['uid']}/review/confirm")
        self.assertTrue(r["reviewed_at"])
        self.assertLess(Record.objects.get(uid=self.co["uid"]).updated_at, since)  # a review is not a change ...
        ch = self.j("get", f"/api/changes?since={since.isoformat().replace('+', '%2B')}")
        self.assertIn(self.co["uid"], [x["uid"] for x in ch["records"]])  # ... but clients still get it
        self.assertEqual(r["reviewed_by"], "read")
        self.assertEqual(len(self.j("get", "/api/bootstrap")["reviews"]), 1)
        self.as_("maintainer")  # someone else's review cannot be confirmed by a read-only user
        self.as_("read")
        self.j("post", f"/api/records/{self.co['uid']}/review/confirm", status=403)

    def test_schedule_rolling_and_campaign(self):
        from datetime import date as d
        from kartei import reviews
        s = Settings.get()
        s.review_mode, s.review_interval_months = "rolling", 12
        s.save()
        self.assertEqual(reviews.schedule(s)["created"], 0)  # all fresh
        Record.objects.filter(pk=self.orphan.pk).update(created_at=timezone.now() - timedelta(days=400))
        Record.objects.filter(uid=self.p["uid"]).update(created_at=timezone.now() - timedelta(days=400))
        self.assertEqual(reviews.schedule(s), {"created": 1, "already_open": 0, "no_responsible": 1})
        Review.objects.update(done_at=timezone.now())
        s.review_mode, s.review_campaign_month, s.review_campaign_day = "campaign", 3, 1
        s.save()
        self.assertIsNone(reviews.schedule(s, today=d(2030, 2, 28)))
        self.assertEqual(reviews.schedule(Settings.get(), today=d(2030, 3, 1))["created"], 2)
        Review.objects.update(done_at=timezone.now())
        self.assertIsNone(reviews.schedule(Settings.get(), today=d(2030, 6, 1)))  # once per year
        self.assertEqual(reviews.schedule(Settings.get(), today=d(2031, 3, 2))["created"], 2)

    def test_change_proposal_accept(self):
        self.as_("admin")
        self.j("post", f"/api/records/{self.p['uid']}/review/request")
        self.as_("read")
        self.j("patch", f"/api/records/{self.p['uid']}", {"version": 1, "last_name": "X"}, status=403)  # can't edit
        bad = self.j("post", f"/api/records/{self.p['uid']}/proposals", {"kind": "change", "data": {"email": "nope"}},
                     status=400)
        self.assertEqual(bad["fields"]["email"], "invalid_email")
        self.j("post", f"/api/records/{self.p['uid']}/proposals", {"kind": "change", "data": {"email": "m@alt.de"}},
               status=400)  # no change
        res = self.j("post", f"/api/records/{self.p['uid']}/proposals",
                     {"kind": "change", "last_name": "Meyer", "data": {"email": "meyer@alt.de", "mobile": ""},
                      "comment": "neue Adresse"}, status=201)["proposal"]
        self.assertEqual(res["changes"], {"core": {"last_name": "Meyer"}, "data": {"email": "meyer@alt.de"}})
        self.assertEqual(res["base"], {"core": {"last_name": "Meier"}, "data": {"email": "m@alt.de"}})
        boot = self.j("get", "/api/bootstrap")
        self.assertEqual(boot["reviews"], [])  # proposing closes the review
        self.assertEqual(len(boot["proposals"]), 1)
        self.assertTrue(Record.objects.get(uid=self.p["uid"]).reviewed_at)
        self.j("post", f"/api/proposals/{res['id']}/accept", status=403)
        self.as_("maintainer")
        self.assertEqual(len(self.j("get", "/api/bootstrap")["proposals"]), 1)
        out = self.j("post", f"/api/proposals/{res['id']}/accept", {})
        self.assertEqual((out["record"]["last_name"], out["record"]["data"]["email"]), ("Meyer", "meyer@alt.de"))
        self.assertEqual(out["proposal"]["status"], "accepted")
        entry = AuditEntry.objects.filter(target_id=self.p["uid"], action="update").first()
        self.assertEqual(entry.changes["proposed_by"], [None, "read"])
        self.assertEqual(self.j("get", "/api/bootstrap")["proposals"], [])
        self.j("post", f"/api/proposals/{res['id']}/accept", status=404)

    def test_proposer_edits_and_withdraws(self):
        self.as_("read")
        p = self.j("post", f"/api/records/{self.p['uid']}/proposals",
                   {"kind": "change", "data": {"email": "x@alt.de"}}, status=201)["proposal"]
        p = self.j("patch", f"/api/proposals/{p['id']}", {"last_name": "Maier", "data": {"email": "y@alt.de"},
                                                          "comment": "Tippfehler"})["proposal"]
        self.assertEqual(p["changes"], {"core": {"last_name": "Maier"}, "data": {"email": "y@alt.de"}})
        self.assertEqual(p["comment"], "Tippfehler")
        self.as_("maintainer")
        self.j("patch", f"/api/proposals/{p['id']}", {"data": {"email": "z@alt.de"}}, status=404)  # not theirs
        self.as_("read")
        self.j("delete", f"/api/proposals/{p['id']}")
        self.assertEqual(self.j("get", "/api/bootstrap")["proposals"], [])
        self.j("delete", f"/api/proposals/{p['id']}", status=404)

    def test_reviewer_adjusts_before_accepting(self):
        self.as_("read")
        p = self.j("post", f"/api/records/{self.p['uid']}/proposals",
                   {"kind": "change", "data": {"email": "mayer@alt.dee"}}, status=201)["proposal"]
        self.as_("maintainer")
        bad = self.j("post", f"/api/proposals/{p['id']}/accept", {"override": {"data": {"email": "kaputt"}}}, status=400)
        self.assertEqual(bad["fields"]["email"], "invalid_email")
        out = self.j("post", f"/api/proposals/{p['id']}/accept", {"override": {"data": {"email": "mayer@alt.de"}}})
        self.assertEqual(out["record"]["data"]["email"], "mayer@alt.de")
        e = AuditEntry.objects.filter(target_id=self.p["uid"], action="update").first()
        self.assertEqual((e.changes["email"][1], e.changes["adjusted"]), ("mayer@alt.de", [None, True]))

    def test_new_contact_proposal(self):
        self.as_("read")
        self.j("post", "/api/records", {"kind": "person", "last_name": "Neu"}, status=403)  # can't create
        self.j("post", "/api/proposals", {"kind": "robot"}, status=400)
        bad = self.j("post", "/api/proposals", {"kind": "person", "data": {"email": "nina@alt.de"}}, status=400)
        self.assertEqual(bad["fields"]["last_name"], "required")
        p = self.j("post", "/api/proposals", {"kind": "person", "first_name": "Nina", "last_name": "Neu",
                                              "company": self.co["uid"], "data": {"email": "nina@alt.de", "internal_contact": "read"},
                                              "comment": "neue Kollegin"}, status=201)["proposal"]
        self.assertEqual((p["kind"], p["record"], p["record_kind"], p["label"]), ("create", None, "person", "Nina Neu"))
        self.assertEqual(p["changes"], {"core": {"first_name": "Nina", "last_name": "Neu", "company": self.co["uid"]},
                                        "data": {"email": "nina@alt.de", "internal_contact": "read"}})
        self.assertEqual([x["id"] for x in self.j("get", "/api/bootstrap")["proposals"]], [p["id"]])
        p = self.j("patch", f"/api/proposals/{p['id']}", {"first_name": "Nina", "last_name": "Neumann",
                                                          "company": self.co["uid"], "data": {"internal_contact": "read"}})["proposal"]
        self.assertEqual(p["label"], "Nina Neumann")
        e = AuditEntry.objects.filter(target_type="proposal", target_id=str(p["id"])).first()
        self.assertEqual((e.action, e.target_label, e.changes["company"]), ("propose", "Nina Neumann", [None, "Alt GmbH"]))
        self.j("post", f"/api/proposals/{p['id']}/accept", status=403)
        self.as_("maintainer")
        out = self.j("post", f"/api/proposals/{p['id']}/accept", {"note": "willkommen"})
        r = out["record"]
        self.assertEqual((r["last_name"], r["company"], out["proposal"]["record"]), ("Neumann", self.co["uid"], r["uid"]))
        created = AuditEntry.objects.get(target_id=r["uid"], action="create")
        self.assertEqual((created.changes["proposed_by"], created.changes["note"]), ([None, "read"], [None, "willkommen"]))
        self.assertEqual(self.j("get", "/api/bootstrap")["proposals"], [])

    def test_new_company_proposal_adjusted_rejected_withdrawn(self):
        self.as_("read")
        new = lambda name: self.j("post", "/api/proposals", {"kind": "company", "name": name, "data": {"internal_contact": "read"}},
                                   status=201)["proposal"]
        a, b, c = new("Neu AG"), new("Unsinn GmbH"), new("Doppelt KG")
        self.j("delete", f"/api/proposals/{c['id']}")
        self.as_("maintainer")
        bad = self.j("post", f"/api/proposals/{a['id']}/accept", {"override": {"name": ""}}, status=400)
        self.assertEqual(bad["fields"]["name"], "required")
        out = self.j("post", f"/api/proposals/{a['id']}/accept", {"override": {"name": "Neu AG & Co. KG", "data": {"internal_contact": "read"}}})
        self.assertEqual(out["record"]["name"], "Neu AG & Co. KG")
        self.assertTrue(AuditEntry.objects.get(target_id=out["record"]["uid"], action="create").changes["adjusted"])
        self.assertIsNone(self.j("post", f"/api/proposals/{b['id']}/reject", {"note": "gibt es nicht"})["record"])
        self.assertTrue(AuditEntry.objects.filter(action="reject", target_type="proposal", target_id=str(b["id"])).exists())
        self.assertEqual(Record.objects.filter(name__in=["Unsinn GmbH", "Doppelt KG"]).count(), 0)
        self.assertEqual(self.j("get", "/api/bootstrap")["proposals"], [])

    def test_delete_proposal_and_reject(self):
        self.as_("read")
        p1 = self.j("post", f"/api/records/{self.co['uid']}/proposals", {"kind": "delete", "comment": "insolvent"},
                    status=201)["proposal"]
        p2 = self.j("post", f"/api/records/{self.p['uid']}/proposals", {"kind": "delete", "reason": "left"},
                    status=201)["proposal"]
        self.as_("maintainer")
        self.j("post", f"/api/proposals/{p2['id']}/reject", {"note": "noch aktiv"})
        self.assertIsNone(Record.objects.get(uid=self.p["uid"]).deleted_at)
        self.j("post", f"/api/proposals/{p1['id']}/accept", {})
        self.assertIsNotNone(Record.objects.get(uid=self.co["uid"]).deleted_at)  # in the recycle bin
        self.assertTrue(AuditEntry.objects.filter(action="reject", target_id=self.p["uid"]).exists())

    def test_delete_reasons(self):
        self.as_("read")
        url = f"/api/records/{self.p['uid']}/proposals"
        bad = self.j("post", url, {"kind": "delete", "reason": "other"}, status=400)
        self.assertEqual(bad["fields"], {"comment": "required"})
        bad = self.j("post", f"/api/records/{self.co['uid']}/proposals", {"kind": "delete", "reason": "left"}, status=400)
        self.assertEqual(bad["fields"], {"reason": "invalid"})
        orphan = f"/api/records/{self.orphan.uid}/proposals"  # no company
        self.j("post", orphan, {"kind": "delete", "reason": "company_gone"}, status=400)

    def test_company_gone_deletes_company_and_persons(self):
        other = self.j("post", "/api/records", {"kind": "person", "last_name": "Zweite", "company": self.co["uid"],
                                                "data": {"internal_contact": "read"}}, status=201)
        self.as_("admin")
        self.j("post", "/api/admin/reviews", {"scope": "all"})
        self.as_("read")
        p = self.j("post", f"/api/records/{self.p['uid']}/proposals", {"kind": "delete", "reason": "company_gone"},
                   status=201)["proposal"]
        self.assertEqual((p["record"], p["reason"]), (self.co["uid"], "company_gone"))  # proposed on the company
        self.assertEqual(self.j("get", "/api/bootstrap")["reviews"], [])  # company and its persons count as answered
        again = self.j("post", f"/api/records/{other['uid']}/proposals", {"kind": "delete", "reason": "company_gone"})
        self.assertEqual((again["existing"], again["proposal"]["id"]), (True, p["id"]))  # no duplicate
        self.as_("maintainer")
        self.j("post", f"/api/proposals/{p['id']}/accept", {})
        gone = Record.objects.filter(uid__in=[self.co["uid"], self.p["uid"], other["uid"]], deleted_at__isnull=False)
        self.assertEqual(gone.count(), 3)
        e = AuditEntry.objects.filter(target_id=self.p["uid"], action="delete").first()
        self.assertEqual((e.changes["reason"], e.changes["proposed_by"]), ([None, "company_gone"], [None, "read"]))

    def test_editing_answers_own_review(self):
        self.as_("admin")
        rec = self.j("get", f"/api/records/{self.p['uid']}")
        self.j("patch", f"/api/records/{self.p['uid']}", {"version": rec["version"], "data": {"internal_contact": "maintainer"}})
        self.j("post", f"/api/records/{self.p['uid']}/review/request")
        self.as_("maintainer")
        rec = self.j("get", f"/api/records/{self.p['uid']}")
        self.j("patch", f"/api/records/{self.p['uid']}", {"version": rec["version"], "data": {"email": "neu@alt.de"}})
        rv = Review.objects.get(record__uid=self.p["uid"], assignee__username="maintainer")
        self.assertEqual(rv.outcome, "edited")
        self.assertEqual(Record.objects.get(uid=self.p["uid"]).reviewed_by.username, "maintainer")

    def test_reassign_and_cancel(self):
        self.as_("admin")
        self.j("post", f"/api/records/{self.p['uid']}/review/request")
        rec = self.j("get", f"/api/records/{self.p['uid']}")
        self.j("patch", f"/api/records/{self.p['uid']}", {"version": rec["version"], "data": {"internal_contact": "maintainer"}})
        rv = Review.objects.get(record__uid=self.p["uid"])
        self.assertEqual((rv.assignee.username, rv.notified_at), ("maintainer", None))
        self.j("delete", f"/api/records/{self.p['uid']}")
        rv.refresh_from_db()
        self.assertEqual(rv.outcome, "cancelled")


class BatchTests(ApiTestBase):
    """Company changes carried over to its persons."""
    def setUp(self):
        super().setUp()
        self.as_("maintainer")
        self.co = self.j("post", "/api/records", {"kind": "company", "name": "Alt GmbH", "data": {
            "company_phone": "0511 1000", "work_street": "Altweg 1", "hotline": "0800 123456"}}, status=201)
        mk = lambda last, data: self.j("post", "/api/records", {"kind": "person", "last_name": last,
                                                                "company": self.co["uid"], "data": data}, status=201)
        self.a = mk("A", {"phone_work": "0511 1000-12", "fax": "0511 1000-99", "work_street": "Altweg 1",
                          "mobile": "0170 1000123"})
        self.b = mk("B", {"phone_work": "+49 511 1000 13", "work_street": "Eigener Weg 5"})  # deviating
        self.changes = {"company_phone": ["0511 1000", "0511 4711"], "work_street": ["Altweg 1", "Neuweg 2"],
                        "hotline": ["0800 123456", "0800 654321"]}

    def test_swap_prefix(self):
        self.assertEqual(batch.swap_prefix("0511 1000-12", "0511 1000", "0511 4711"), "0511 4711-12")
        self.assertEqual(batch.swap_prefix("(0511) 1000 12", "05111000", "0511 4711"), "0511 4711 12")
        self.assertIsNone(batch.swap_prefix("0511 1001-12", "0511 1000", "0511 4711"))
        self.assertIsNone(batch.swap_prefix("110", "110", "112"))  # too short to be safe
        self.assertEqual(batch.swap_prefix("+49 511 1000 13", "0511 1000", "0511 4711"), "0511 4711 13")
        self.assertEqual(batch.swap_prefix("0049-511-1000-14", "0511 1000", "0511 4711"), "0511 4711-14")

    def test_preview_and_apply(self):
        items = self.j("post", f"/api/records/{self.co['uid']}/propagate", {"changes": self.changes})["items"]
        got = {(i["name"], i["key"]): i["new"] for i in items}
        self.assertEqual(got, {("A", "phone_work"): "+49 511 4711-12", ("A", "fax"): "+49 511 4711-99",
                               ("A", "work_street"): "Neuweg 2", ("B", "phone_work"): "+49 511 4711 13"})  # B keeps its street
        sel = [{"uid": self.a["uid"], "key": "phone_work"}, {"uid": self.a["uid"], "key": "work_street"}]
        res = self.j("post", f"/api/records/{self.co['uid']}/propagate", {"changes": self.changes, "selected": sel})
        self.assertEqual(res, {"updated": 1, "errors": []})
        a = Record.objects.get(uid=self.a["uid"])
        self.assertEqual((a.data["phone_work"], a.data["fax"], a.data["work_street"], a.version),
                         ("+49 511 4711-12", "+49 511 1000-99", "Neuweg 2", self.a["version"] + 1))
        e = AuditEntry.objects.filter(target_id=self.a["uid"], action="update").first()
        self.assertEqual(e.changes["batch_from"], [None, "Alt GmbH"])
        self.as_("read")
        self.j("post", f"/api/records/{self.co['uid']}/propagate", {"changes": self.changes}, status=403)


class CardDavTests(TestCase):
    def setUp(self):
        self.person, self.company = sample_records()
        s = Settings.get()
        s.carddav_enabled, s.carddav_password_hash = True, make_password("phone-secret")
        s.save()
        self.auth = "Basic " + base64.b64encode(b"carddav:phone-secret").decode()

    def dav(self, method, path, body="", **headers):
        return self.client.generic(method, path, body, content_type="application/xml",
                                   HTTP_AUTHORIZATION=self.auth, **headers)

    def test_auth_and_readonly(self):
        self.assertEqual(self.client.generic("PROPFIND", "/carddav/").status_code, 401)
        self.assertEqual(self.dav("PUT", f"/carddav/kontakte/{PERSON_UID}.vcf", "x").status_code, 403)
        self.assertEqual(self.dav("DELETE", f"/carddav/kontakte/{PERSON_UID}.vcf").status_code, 403)

    def sync(self, token=""):
        body = ('<d:sync-collection xmlns:d="DAV:" xmlns:c="urn:ietf:params:xml:ns:carddav"><d:sync-token>' + token +
                '</d:sync-token><d:sync-level>1</d:sync-level><d:prop><d:getetag/></d:prop></d:sync-collection>')
        return self.dav("REPORT", "/carddav/kontakte/", body)

    def test_sync_collection(self):
        r = self.dav("PROPFIND", "/carddav/kontakte/", HTTP_DEPTH="0")
        self.assertIn(b"<d:sync-collection />", r.content.replace(b"<d:sync-collection/>", b"<d:sync-collection />"))
        self.assertIn(b"<d:sync-token>urn:kartei:sync:", r.content)
        r = self.sync()
        self.assertEqual((r.status_code, r.content.count(b".vcf</d:href>")), (207, 3))
        Record.objects.update(updated_at=timezone.now() - timedelta(hours=1))
        token = re.findall(rb"<d:sync-token>([^<]+)</d:sync-token>", r.content)[-1].decode()
        self.assertEqual(self.sync(token).content.count(b"<d:response>"), 0)
        Record.objects.filter(uid=self.company.uid).update(version=2, updated_at=timezone.now())  # person's card changes too
        Record.objects.filter(uid=PERSON_UID).update(deleted_at=timezone.now())
        r = self.sync(token)
        self.assertIn(f"{self.company.uid}.vcf</d:href><d:propstat>".encode(), r.content)
        self.assertIn(f"{PERSON_UID}.vcf</d:href><d:status>HTTP/1.1 404 Not Found".encode(), r.content)
        # a changed vCard mapping or purged contacts: the token is no longer usable, the phone starts over
        s = Settings.get()
        s.base_url = "https://kartei.example.org"
        s.save()
        self.assertEqual(self.sync(token).status_code, 403)
        token = re.findall(rb"<d:sync-token>([^<]+)</d:sync-token>", self.sync().content)[-1].decode()
        self.assertEqual(self.sync(token).status_code, 207)
        AuditEntry.objects.create(action="purge", target_type="record", target_id=PERSON_UID)
        r = self.sync(token)
        self.assertEqual(r.status_code, 403)
        self.assertIn(b"valid-sync-token", r.content)

    def test_new_password_ends_cached_login(self):
        depth0 = lambda: self.dav("PROPFIND", "/carddav/kontakte/", HTTP_DEPTH="0").status_code
        self.assertEqual(depth0(), 207)  # the successful check is now cached
        Settings.objects.filter(pk=1).update(carddav_password_hash=make_password("new-secret"))
        self.assertEqual(depth0(), 401)

    def test_discovery_and_sync(self):
        r = self.dav("PROPFIND", "/carddav/", HTTP_DEPTH="0", body=
                     '<d:propfind xmlns:d="DAV:" xmlns:c="urn:ietf:params:xml:ns:carddav"><d:prop>'
                     '<d:current-user-principal/><c:addressbook-home-set/></d:prop></d:propfind>')
        self.assertEqual(r.status_code, 207)
        self.assertIn(b"/carddav/", r.content)
        r = self.dav("PROPFIND", "/carddav/kontakte/", HTTP_DEPTH="1")
        self.assertEqual(r.content.count(b".vcf</d:href>"), 3)
        self.assertIn(b"getctag", r.content)
        body = ('<c:addressbook-multiget xmlns:d="DAV:" xmlns:c="urn:ietf:params:xml:ns:carddav"><d:prop>'
                f'<d:getetag/><c:address-data/></d:prop><d:href>/carddav/kontakte/{PERSON_UID}.vcf</d:href>'
                '</c:addressbook-multiget>')
        r = self.dav("REPORT", "/carddav/kontakte/", body, HTTP_USER_AGENT="iOS/17.0 (21A329) dataaccessd/1.0")
        self.assertEqual(r.status_code, 207)
        self.assertIn("ORG:Muster IT GmbH;Informatik".encode(), r.content)
        self.assertEqual(r.content.count(b"BEGIN:VCARD"), 1)
        r = self.dav("GET", f"/carddav/kontakte/{PERSON_UID}.vcf", HTTP_USER_AGENT="DAVx5/4.4")
        self.assertIn(b"ROLE:Informatik", r.content)

    def test_restricted_fields_hidden(self):
        f = FieldDef.objects.get(key="description")
        f.visible_to = "maintainer"
        f.save()
        r = self.dav("GET", f"/carddav/kontakte/{PERSON_UID}.vcf")
        self.assertNotIn(b"Kaffe", r.content)


class ExportTests(TestCase):
    def setUp(self):
        sample_records()
        self.dir = tempfile.mkdtemp()
        patcher = mock.patch.object(exports, "EXPORT_ROOT", self.dir)
        patcher.start()
        self.addCleanup(patcher.stop)

    def profile(self, fmt="csv", **kw):
        cols = [{"header": h, "person": p, "company": c} for h, p, c in EXPORT_COLUMNS]
        return ExportProfile.objects.create(name="t", format=fmt, path=f"{self.dir}/out.{fmt}", columns=cols, **kw)

    def test_csv_layout(self):
        p = self.profile(options={"delimiter": ",", "encoding": "utf-8", "line_ending": "lf"})
        self.assertEqual(exports.run(p), "ok", p.last_message)
        text = open(p.path, encoding="utf-8").read()
        header, *rows = text.strip().split("\n")
        self.assertEqual(header.split(","), [h for h, _, _ in EXPORT_COLUMNS])
        self.assertIn("ACME,,,,ACME GmbH", text)  # company: firstname = short name
        self.assertEqual(exports.run(p), "skipped")
        self.assertEqual(os.stat(p.path).st_mode & 0o777, 0o644)

    def test_multiline_and_cp1252(self):
        Record.objects.filter(kind="company", name="ACME GmbH").update(data={"description": "Zeile 1\nZeile 2 €"})
        p = self.profile(options={"encoding": "cp1252", "newlines": "keep"})
        exports.run(p)
        import csv
        with open(p.path, encoding="cp1252", newline="") as fh:
            rows = list(csv.reader(fh, delimiter=";"))
        self.assertIn("Zeile 1\nZeile 2 €", [r[-1] for r in rows])

    def test_guard(self):
        p = self.profile()
        self.assertEqual(exports.run(p), "ok")
        Record.objects.filter(kind="person").update(deleted_at="2026-01-01T00:00Z")
        Record.objects.filter(name="Muster IT GmbH").update(deleted_at="2026-01-01T00:00Z")
        before = open(p.path, "rb").read()
        self.assertEqual(exports.run(p), "error")
        self.assertIn("@export_dropped", p.last_message)
        self.assertEqual(open(p.path, "rb").read(), before)  # file untouched
        Record.objects.update(deleted_at="2026-01-01T00:00Z")
        self.assertEqual(exports.run(p), "error")
        self.assertIn("@export_empty", p.last_message)
        self.assertEqual(exports.run(p, force=True), "ok")

    def test_path_outside_root_rejected(self):
        p = self.profile()
        p.path = "/etc/passwd.csv"
        self.assertEqual(exports.run(p), "error")
        self.assertTrue(os.path.exists("/etc/passwd"))

    def test_xlsx(self):
        from openpyxl import load_workbook
        p = self.profile("xlsx")
        self.assertEqual(exports.run(p), "ok", p.last_message)
        ws = load_workbook(p.path).active
        self.assertEqual(ws["A1"].value, "ID")
        self.assertEqual(ws.freeze_panes, "A2")


class ImportTests(TestCase):
    CSV = ("ID;Anrede;Vorname;Nachname;Funktion;Abteilung;Firma;PLZ;Telefon;E-Mail;Schlagworte;Notizen\n"
           "F-5;;ACME;;;;ACME GmbH;10115;030 123;info@acme.com;Kunde |##| Partner;\"zwei\nZeilen\"\n"
           "P-7;Herr;Max;Muster;CEO;IT;ACME GmbH;10115;030 124;max@acme.com;Kunde;\n"
           "P-8;;Ohne;Firma;;;Neue Firma;;kaputt!;;;\n").encode("cp1252")

    def setUp(self):
        load_fields()

    def test_import_idempotent(self):
        dry = import_test_csv(self.CSV, dry_run=True)
        self.assertEqual(dry["persons_created"], 2)
        self.assertFalse(Record.objects.exists())
        res = import_test_csv(self.CSV, dry_run=False)
        self.assertEqual((res["companies_created"], res["persons_created"]), (2, 2))
        self.assertTrue(any("invalid_phone" in w for w in res["warnings"]))
        max_ = Record.objects.get(legacy_id="P-7")
        self.assertEqual(max_.company.legacy_id, "F-5")
        self.assertEqual(max_.data["salutation"], "Herr")
        acme = Record.objects.get(legacy_id="F-5")
        self.assertEqual(acme.data["keywords"], ["Kunde", "Partner"])
        self.assertEqual(acme.data["description"], "zwei\nZeilen")
        again = import_test_csv(self.CSV, dry_run=False)
        self.assertEqual(again["persons_updated"] + again["persons_created"], 0)


class FakeGraph:
    """In-memory stand-in for graph.http: one group, mailboxes with contact folders."""
    def __init__(self, users):
        self.users = dict(users)  # id -> mail
        self.folders = {}  # (user, fid) -> {"name", "contacts": {cid: body}}
        self.calls = 0
        self.n = 0
        self.photo_fails = False

    def _id(self):
        self.n += 1
        return f"id{self.n}"

    def __call__(self, method, url, body=None, token=None, form=None, ctype=None):
        from urllib.parse import unquote
        self.calls += 1
        if "login.microsoftonline.com" in url:
            return {"access_token": "tok"}
        path = unquote(url.removeprefix(graph.GRAPH))
        if path == "/$batch":
            return {"responses": [{"id": r["id"], **self._one(r["method"], r["url"], r.get("body"))}
                                  for r in body["requests"]]}
        res = self._one(method, path, body)
        if res["status"] >= 400:
            raise graph.GraphError(f"HTTP {res['status']}")
        return res["body"]

    def _one(self, method, path, body):
        p, _, q = path.partition("?")
        parts = p.strip("/").split("/")
        if parts[0] == "groups":
            return {"status": 200, "body": {"value": [{"@odata.type": "#microsoft.graph.user", "id": u, "mail": m}
                                                      for u, m in self.users.items()]}}
        user = parts[1]
        if user not in self.users and not any(k[0] == user for k in self.folders):
            return {"status": 404, "body": {"error": {"message": "no mailbox"}}}
        if len(parts) == 3:  # contactFolders
            if method == "POST":
                fid = self._id()
                self.folders[(user, fid)] = {"name": body["displayName"], "contacts": {}}
                return {"status": 201, "body": {"id": fid}}
            name = q.split("'")[1]
            return {"status": 200, "body": {"value": [{"id": f} for (u, f), d in self.folders.items()
                                                      if u == user and d["name"] == name]}}
        folder = self.folders.get((user, parts[3]))
        if folder is None:
            return {"status": 404, "body": {}}
        if len(parts) == 4 and method == "DELETE":
            del self.folders[(user, parts[3])]
            return {"status": 204, "body": {}}
        contacts = folder["contacts"]
        if len(parts) == 5 and method == "GET":
            return {"status": 200, "body": {"value": [
                {"id": cid, **({"singleValueExtendedProperties": c["singleValueExtendedProperties"]}
                               if "singleValueExtendedProperties" in c else {})} for cid, c in contacts.items()]}}
        if method == "POST":
            cid = self._id()
            contacts[cid] = body
            return {"status": 201, "body": {"id": cid}}
        if parts[5] not in contacts:
            return {"status": 404, "body": {}}
        if parts[6:] == ["photo", "$value"]:
            if self.photo_fails:
                return {"status": 500, "body": {"error": {"message": "boom"}}}
            contacts[parts[5]]["photo"] = body
            return {"status": 200, "body": {}}
        if method == "PATCH":
            contacts[parts[5]] |= body
        else:
            del contacts[parts[5]]
        return {"status": 200, "body": {}}

    def contacts(self, user):
        return [c for (u, _), d in self.folders.items() if u == user for c in d["contacts"].values()]


class GraphSyncTests(TestCase):
    def setUp(self):
        self.person, self.company = sample_records()
        s = Settings.get()
        s.graph_enabled, s.graph_tenant_id, s.graph_client_id, s.graph_group_id = True, "t", "c", "g"
        s.graph_client_secret = "secret"
        s.save()
        self.fake = FakeGraph({"u1": "a@x.de", "u2": "b@x.de"})
        patcher = mock.patch.object(graph, "http", self.fake)
        patcher.start()
        self.addCleanup(patcher.stop)

    def by_name(self, user, name):
        return next(c for c in self.fake.contacts(user) if c["displayName"] == name)

    def test_mapping(self):
        c = graph.contact(self.person, list(FieldDef.objects.all()))
        self.assertEqual((c["givenName"], c["surname"], c["companyName"]), ("Alexander", "Meier", "Muster IT GmbH"))
        self.assertEqual(c["mobilePhone"], "01701234567")
        self.assertEqual(c["businessPhones"], ["05111234567"])
        self.assertEqual(c["homePhones"], ["050311234567"])
        self.assertEqual(c["businessAddress"], {"street": "Empfang\nMuster Straße 1", "city": "Hannover", "state": "",
                                                "postalCode": "30159", "countryOrRegion": "Deutschland"})
        self.assertEqual(c["birthday"], "1994-01-01T00:00:00Z")
        self.assertEqual(len(c["emailAddresses"]), 3)
        self.assertIn("Fax: 051112345671117", c["personalNotes"])
        self.assertEqual(c["department"], "Informatik")
        co = graph.contact(self.company, list(FieldDef.objects.all()))
        self.assertEqual((co["displayName"], co["companyName"], co["givenName"]), ("ACME GmbH", "ACME GmbH", ""))

    def test_sync_create_update_delete(self):
        st = graph.run()
        self.assertEqual((st["mailboxes"], st["created"], st["failed"]), (2, 6, []))
        self.assertEqual(len(self.fake.contacts("u1")), 3)
        # unchanged: no writes
        before = self.fake.calls
        st = graph.run()
        self.assertEqual((st["created"], st["updated"], st["deleted"]), (0, 0, 0))
        self.assertLess(self.fake.calls - before, 10)
        # edit -> PATCH, foreign contact on the phone -> removed, record deleted -> removed
        self.person.first_name = "Alex"
        self.person.save()
        folder = next(d for (u, _), d in self.fake.folders.items() if u == "u1")
        folder["contacts"]["foreign"] = {"displayName": "added on phone"}
        Record.objects.filter(pk=self.person.company_id).update(deleted_at=timezone.now())
        st = graph.run()
        self.assertEqual((st["updated"], st["deleted"]), (2, 3))
        self.assertEqual(sorted(c["displayName"] for c in self.fake.contacts("u1")), ["ACME GmbH", "Alex Meier"])
        self.assertEqual(self.by_name("u1", "Alex Meier")["companyName"], "")

    def test_leaver_folder_removed_and_guards(self):
        graph.run()
        del self.fake.users["u2"]
        st = graph.run()
        self.assertEqual(st["mailboxes"], 1)
        self.assertEqual(self.fake.contacts("u2"), [])
        self.assertEqual(Settings.get().graph_mailboxes, ["u1"])
        Record.objects.update(deleted_at=timezone.now())
        st = graph.run()
        self.assertIn("@graph_no_records", st["error"])
        self.assertEqual(len(self.fake.contacts("u1")), 3)  # untouched

    def test_personal_folders(self):
        FieldDef.objects.create(key="owner", label_de="Zuständig", label_en="Owner", type="user", applies_to="both")
        FieldDef.objects.create(key="internal", label_de="Intern", label_en="Internal", type="users", applies_to="both")
        User.objects.create(username="anna", email="A@x.de")  # mailbox u1 by mail
        User.objects.create(username="b")  # mailbox u2 by the local part
        Record.objects.filter(pk=self.person.pk).update(data={**self.person.data, "owner": "anna", "internal": ["b"]})
        Record.objects.filter(pk=self.company.pk).update(data={**self.company.data, "internal": ["anna", "b"]})
        s = Settings.get()
        s.graph_folder_responsible, s.graph_folder_internal = "Meine", "Intern"
        s.save()
        self.assertEqual(graph.run()["failed"], [])
        folder = lambda u, n: sorted(c["displayName"] for (uu, _), d in self.fake.folders.items()
                                     if uu == u and d["name"] == n for c in d["contacts"].values())
        self.assertEqual(len(folder("u1", "Kartei")), 3)
        self.assertEqual(folder("u1", "Meine"), ["Alexander Meier"])
        self.assertEqual(folder("u1", "Intern"), ["ACME GmbH"])
        self.assertEqual(folder("u2", "Meine"), [])
        self.assertEqual(folder("u2", "Intern"), ["ACME GmbH", "Alexander Meier"])
        Record.objects.filter(pk=self.company.pk).update(data={})  # no longer an internal contact
        graph.run()
        self.assertEqual(folder("u1", "Intern"), [])
        del self.fake.users["u1"]  # leaver: all three folders go
        graph.run()
        self.assertEqual([d["name"] for (u, _), d in self.fake.folders.items() if u == "u1"], [])

    def test_photos(self):
        jpeg = b"\xff\xd8\xff" + b"x" * 50
        Record.objects.filter(pk=self.person.pk).update(image=jpeg, image_type="jpeg")
        self.fake.photo_fails = True
        st = graph.run()
        self.assertEqual(st["photos"], 0)
        self.assertEqual(len(st["failed"]), 2)
        self.fake.photo_fails = False
        st = graph.run()  # retried, data unchanged
        self.assertEqual((st["photos"], st["created"], st["updated"], st["failed"]), (2, 0, 0, []))
        self.assertEqual(self.by_name("u1", "Alexander Meier")["photo"], jpeg)
        self.assertEqual(graph.run()["photos"], 0)  # nothing to do
        Record.objects.filter(pk=self.person.pk).update(image=None, image_type="")
        st = graph.run()  # removal: contact recreated without photo
        self.assertEqual(st["created"], 2)
        self.assertNotIn("photo", self.by_name("u1", "Alexander Meier"))
        self.assertEqual(len(self.fake.contacts("u1")), 3)

    def test_sync_now(self):
        Settings.objects.update(graph_sync_requested=timezone.now())
        graph.run()
        self.assertIsNone(Settings.get().graph_sync_requested)

    def test_disabled_and_secret_masked(self):
        self.assertEqual(Settings.get().graph_client_secret, "secret")
        self.assertNotIn("secret", Settings.get().graph_client_secret_enc)
        self.assertEqual(graph.test(Settings.get())["mailboxes"], 2)
        Settings.objects.update(graph_group_id="")
        self.assertIn("@graph_no_group", graph.test(Settings.get())["error"])
        Settings.objects.update(graph_client_id="")
        self.assertIn("@graph_not_configured", graph.test(Settings.get())["error"])
        Settings.objects.update(graph_enabled=False)
        self.assertIsNone(graph.run())


class PwaTests(TestCase):
    def test_service_worker_from_root(self):
        from django.conf import settings as dj
        with tempfile.TemporaryDirectory() as d, mock.patch.object(dj, "FRONTEND_DIST", __import__("pathlib").Path(d)):
            self.assertEqual(self.client.get("/sw.js").status_code, 404)
            open(f"{d}/sw.js", "w").write("// sw")
            r = self.client.get("/sw.js")
            self.assertEqual((r.status_code, r["Content-Type"]), (200, "text/javascript"))
            self.assertIn("no-cache", r["Cache-Control"])


class RenameTests(TestCase):
    def test_old_crm_links_redirect(self):
        r = self.client.get("/crm/person/ad081c47-96f6-4b88-9d9b-67faf688772d?x=1")
        self.assertEqual((r.status_code, r["Location"]), (301, "/person/ad081c47-96f6-4b88-9d9b-67faf688772d?x=1"))
        self.assertEqual(self.client.get("/crm")["Location"], "/")


class OpsTests(TestCase):
    def test_client_ip_behind_proxy(self):
        from django.test import RequestFactory, override_settings
        from .net import client_ip
        rf = RequestFactory()
        req = rf.get("/", REMOTE_ADDR="172.18.0.1", HTTP_X_FORWARDED_FOR="6.6.6.6, 10.0.0.7")
        self.assertEqual(client_ip(req), "10.0.0.7")  # what the proxy appended, not the client's claim
        with override_settings(PROXY_HEADERS=False):
            self.assertEqual(client_ip(req), "172.18.0.1")
        self.assertEqual(client_ip(rf.get("/", REMOTE_ADDR="10.0.0.8")), "10.0.0.8")

    def test_init_app_creates_admin_only_on_fresh_install(self):
        from django.core.management import call_command
        with mock.patch.dict(os.environ, {"ADMIN_USERNAME": "boss", "ADMIN_PASSWORD": "pw-boss-12345"}):
            call_command("init_app", stdout=open(os.devnull, "w"))
            self.assertTrue(User.objects.get(username="boss").is_superuser)
            User.objects.filter(username="boss").update(is_superuser=False)  # an LDAP admin exists instead
            Profile.objects.filter(user__username="boss").update(role="admin")
            call_command("init_app", stdout=open(os.devnull, "w"))  # must not crash or add a second admin
        self.assertEqual(User.objects.count(), 1)


class FieldFileTests(ApiTestBase):
    """Field definitions as a JSON file: export, preview, apply."""
    def test_roundtrip_matches_example(self):
        self.as_("admin")
        FieldDef.objects.filter(key="internal_contact").update(required=True)  # as in the file
        doc = self.j("get", "/api/admin/fields/export")
        self.assertEqual(doc, example("fields-example.json"))
        res = self.j("post", "/api/admin/fields/import", {"document": doc})["fields"]
        self.assertEqual({r["status"] for r in res}, {"unchanged"})

    def test_preview_and_apply(self):
        self.as_("admin")
        Record.objects.create(kind="person", last_name="A", data={"email": "a@x.de"})
        doc = example("fields-example.json")
        fields = {f["key"]: f for f in doc["fields"]}
        fields["mobile"]["label"]["de"] = "Handy"
        fields["mobile"]["import_aliases"] = "Handynummer, Mobil"
        fields["email"]["type"] = "integer"  # has data: incompatible
        doc["fields"] = [fields["mobile"], fields["email"], {"key": "Bad Key"},
                         {**fields["fax"], "key": "fax2", "copy_from": "company_fax"},
                         {**fields["fax"], "key": "fax3", "copy_from": "nope"}]
        res = {r["key"]: r for r in self.j("post", "/api/admin/fields/import", {"document": doc})["fields"]}
        self.assertEqual((res["mobile"]["status"], res["mobile"]["changes"]), ("changed", ["import_aliases", "label_de"]))
        self.assertEqual((res["email"]["status"], res["email"]["errors"]), ("invalid", {"type": "incompatible_type"}))
        self.assertEqual(res["Bad Key"]["status"], "invalid")
        self.assertEqual((res["fax2"]["status"], res["fax3"]["errors"]), ("new", {"copy_from": "invalid"}))
        self.assertFalse(FieldDef.objects.filter(key="fax2").exists())  # preview only
        self.j("post", "/api/admin/fields/import", {"document": doc, "apply": True})
        m = FieldDef.objects.get(key="mobile")
        self.assertEqual((m.label_de, m.import_aliases), ("Handy", ["Handynummer", "Mobil"]))
        self.assertEqual(FieldDef.objects.get(key="fax2").copy_from.key, "company_fax")
        self.assertEqual(FieldDef.objects.get(key="email").type, "email")
        self.assertTrue(FieldDef.objects.filter(key="website").exists())  # not in the file: untouched
        self.assertTrue(AuditEntry.objects.filter(target_type="field", action="update", target_label="Handy").exists())
        self.j("post", "/api/admin/fields/import", {"document": {"fields": []}}, status=400)

    def test_init_app_loads_initial_fields(self):
        from django.core.management import call_command
        FieldDef.objects.all().delete()
        Settings.objects.all().delete()
        call_command("init_app", stdout=open(os.devnull, "w"))
        self.assertEqual(list(FieldDef.objects.values_list("key", flat=True)),
                         [f["key"] for f in example("fields-default.json")["fields"]])
        FieldDef.objects.all().delete()
        Settings.objects.all().delete()
        with mock.patch.dict(os.environ, {"INITIAL_FIELDS": ""}):
            call_command("init_app", stdout=open(os.devnull, "w"))
        self.assertFalse(FieldDef.objects.exists())


class CsvMappingTests(ApiTestBase):
    """Generic CSV: suggested mapping, kind column, new fields."""
    CSV = ("Art;Vorname;Nachname;Firma;E-Mail;Telefon Zentrale;Lieblingsfarbe;Notiz\n"
           "Firma;;;Ärzte AG;;0511 1;;x\n"
           "Person;Eva;Huber;Ärzte AG;eva@aerzte.de;;blau;y\n"
           "Kontakt;Tom;;;;;grün;\n"
           "Alien;Ute;Ufo;;;;;\n").encode()

    def upload(self, url, **data):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return self.client.post(url, {"file": SimpleUploadedFile("x.csv", self.CSV), **data})

    def test_preview_suggests(self):
        self.as_("admin")
        r = self.upload("/api/admin/import/preview").json()
        cols = r["mapping"]["columns"]
        self.assertEqual(r["mapping"]["kind"]["column"], "Art")
        self.assertEqual(cols["Vorname"], {"person": "core:first_name", "company": ""})
        self.assertEqual(cols["Firma"], {"person": "core:company", "company": "core:name"})
        self.assertEqual(cols["E-Mail"], {"person": "field:email", "company": "field:company_email"})  # label / alias
        self.assertEqual(cols["Telefon Zentrale"], {"person": "", "company": "field:company_phone"})
        self.assertEqual(cols["Lieblingsfarbe"], {"person": "", "company": ""})
        self.assertEqual((r["rows"], r["samples"][1][1]), (4, "Eva"))

    def test_import_with_mapping(self):
        self.as_("admin")
        mapping = {"columns": {"Lieblingsfarbe": {"person": "new", "company": ""},
                               "Notiz": {"person": "field:description", "company": "field:description"}}}
        dry = self.upload("/api/admin/import", dry_run="1", mapping=json.dumps(mapping)).json()
        self.assertEqual((dry["companies_created"], dry["persons_created"]), (1, 3))
        self.assertFalse(FieldDef.objects.filter(key="lieblingsfarbe").exists())
        res = self.upload("/api/admin/import", dry_run="0", mapping=json.dumps(mapping)).json()
        self.assertTrue(any(w.startswith("@import_kind_unknown") and "Alien" in w for w in res["warnings"]))
        eva = Record.objects.get(first_name="Eva")
        self.assertEqual((eva.company.name, eva.data["lieblingsfarbe"], eva.data["description"]), ("Ärzte AG", "blau", "y"))
        self.assertEqual(Record.objects.get(kind="company").data, {"company_phone": "+49 511 1", "description": "x"})
        self.assertEqual(Record.objects.get(first_name="Tom").kind, "person")
        self.assertEqual(FieldDef.objects.get(key="lieblingsfarbe").applies_to, "person")
        bad = {"columns": {"Notiz": {"person": "field:company_phone", "company": ""}}}
        r = self.upload("/api/admin/import", dry_run="1", mapping=json.dumps(bad))
        self.assertEqual(r.status_code, 400)
        self.assertIn("@import_bad_target", r.json()["detail"])


class SeedDemoTests(TestCase):
    def test_seed_demo_once(self):
        import io
        from django.core.management import CommandError, call_command
        out = io.StringIO()
        call_command("seed_demo", password="demo-pass-123", stdout=out)
        self.assertGreater(Record.objects.filter(kind="person").count(), 30)
        self.assertEqual(Record.objects.filter(kind="company").count(), 6)
        self.assertTrue(User.objects.get(username="kkramer").check_password("demo-pass-123"))
        self.assertTrue(Review.objects.exists())
        self.assertEqual(Proposal.objects.filter(status="pending").count(), 4)  # incl. one new person
        new = Proposal.objects.get(kind="create")
        self.assertEqual(new.changes["data"]["job_title"], "Einkauf")  # field values survive the seeding
        self.assertEqual(Record.objects.filter(deleted_at__isnull=False).count(), 3)  # recycle bin has content
        kw = FieldDef.objects.get(key="keywords")  # used keywords are pickable options
        self.assertIn("Kunde", [o["value"] for o in kw.options])
        with self.assertRaises(CommandError):  # never on top of real data
            call_command("seed_demo", stdout=io.StringIO())


class CsvReadTests(ApiTestBase):
    """Files as real exports produce them: none of these may end in a server error."""
    def preview(self, data, status=200):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.as_("admin")
        r = self.client.post("/api/admin/import/preview", {"file": SimpleUploadedFile("k.csv", data, content_type="text/csv")})
        self.assertEqual(r.status_code, status, r.content[:300])
        return r.json()

    def test_line_endings_encodings_and_long_fields(self):
        base = "Vorname;Nachname;Firma\nJürgen;Groß;Müller GmbH\nAnna;Bär;\n"
        for data in (base.replace("\n", "\r").encode("cp1252"),  # "CSV (Macintosh)"
                     base.replace("\n", "\r\n").encode("utf-8-sig"), base.encode("utf-16"), base.encode("cp1252")):
            p = self.preview(data)
            self.assertEqual(p["headers"], ["Vorname", "Nachname", "Firma"])
            self.assertEqual(p["samples"][0][:2], ["Jürgen", "Groß"])
        self.assertEqual(self.preview(("A;B\nx;" + "y" * 200000 + "\n").encode())["headers"], ["A", "B"])

    def test_excel_and_broken_files_get_a_message(self):
        self.assertIn("import_excel", self.preview(b"PK\x03\x04rest", status=400)["detail"])
        self.assertIn("import_unreadable", self.preview(b'A;B\n"x;y\n' + b"z" * 20 * 1024 * 1024, status=400)["detail"])

    def test_large_upload(self):
        rows = "".join(f"Max{i};Muster{i};m{i}@example.com;{'x' * 120}\n" for i in range(80000))  # ~12 MB
        self.assertEqual(len(self.preview(("Vorname;Nachname;E-Mail;Notiz\n" + rows).encode())["headers"]), 4)
