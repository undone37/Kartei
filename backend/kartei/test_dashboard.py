import csv
import io
import uuid

from openpyxl import load_workbook

from .models import AuditEntry, FieldDef, Profile, Record, SavedFilter
from .tests import ApiTestBase


class RecentFavoriteTests(ApiTestBase):
    def test_recent_views(self):
        recs = [Record.objects.create(kind="company", name=f"C{i}") for i in range(25)]
        self.as_("read")
        for r in recs:
            self.j("post", "/api/me/viewed", {"uid": str(r.uid)})
        out = self.j("post", "/api/me/viewed", {"uid": str(recs[3].uid)})["recent"]
        self.assertEqual(len(out), 20)
        self.assertEqual(out[:2], [str(recs[3].uid), str(recs[24].uid)])
        self.assertEqual(out.count(str(recs[3].uid)), 1)
        boot = self.j("get", "/api/bootstrap")
        self.assertEqual(boot["recent"], out)
        ch = self.j("get", f"/api/changes?since={boot['cursor'].replace('+', '%2B')}")
        self.assertEqual(ch["recent"], out)
        self.j("post", "/api/me/viewed", {"uid": str(uuid.uuid4())}, status=404)
        self.j("post", "/api/me/viewed", {"uid": "nope"}, status=404)

    def test_favorites(self):
        a, b = Record.objects.create(kind="company", name="A"), Record.objects.create(kind="company", name="B")
        self.as_("read")
        self.j("post", "/api/me/favorite", {"uid": str(a.uid), "on": True})
        self.assertEqual(self.j("post", "/api/me/favorite", {"uid": str(b.uid), "on": True})["favorites"],
                         [str(b.uid), str(a.uid)])
        self.assertEqual(self.j("post", "/api/me/favorite", {"uid": str(b.uid), "on": False})["favorites"], [str(a.uid)])
        self.assertEqual(self.j("get", "/api/bootstrap")["favorites"], [str(a.uid)])
        self.assertEqual(Profile.objects.get(user=self.users["maintainer"]).favorites, [])  # per user


class SavedFilterTests(ApiTestBase):
    def test_lifecycle_and_sharing(self):
        self.as_("read")
        state = {"q": "acme", "scope": "responsible", "kind": "person", "filters": {"keywords": ["A"], "x": []}, "junk": 1}
        f = self.j("post", "/api/filters", {"name": "  Meine  Firmen ", "state": state}, status=201)
        self.assertEqual(f["name"], "Meine Firmen")
        self.assertEqual(f["state"], {"q": "acme", "scope": "responsible", "kind": "person", "filters": {"keywords": ["A"]}})
        self.assertTrue(f["mine"])
        bad = self.j("post", "/api/filters", {"name": "x", "shared": True}, status=400)
        self.assertEqual(bad["fields"], {"shared": "forbidden"})
        self.j("post", "/api/filters", {"name": "", "state": state}, status=400)
        self.j("post", "/api/filters", {"name": "x", "state": {"filters": {"k": "notalist"}}}, status=400)

        self.as_("maintainer")
        shared = self.j("post", "/api/filters", {"name": "Alle Kunden", "state": {}, "shared": True}, status=201)
        self.assertEqual([x["name"] for x in self.j("get", "/api/bootstrap")["filters"]], ["Alle Kunden"])
        self.j("patch", f"/api/filters/{f['id']}", {"name": "mine now"}, status=404)  # not the owner
        self.j("delete", f"/api/filters/{f['id']}", status=404)

        self.as_("read")
        boot = self.j("get", "/api/bootstrap")["filters"]
        self.assertEqual({(x["name"], x["mine"]) for x in boot}, {("Alle Kunden", False), ("Meine Firmen", True)})
        self.j("patch", f"/api/filters/{shared['id']}", {"name": "x"}, status=404)
        self.assertEqual(self.j("patch", f"/api/filters/{f['id']}", {"name": "Umbenannt"})["name"], "Umbenannt")
        self.j("delete", f"/api/filters/{f['id']}")
        self.assertFalse(SavedFilter.objects.filter(pk=f["id"]).exists())

    def test_advanced_filter_state(self):
        self.as_("read")
        adv = {"mode": "any", "junk": 1, "items": [
            {"field": "@reviewed_at", "op": "older", "value": 365, "x": 1},
            {"mode": "none", "items": [{"field": "keywords", "op": "any", "value": ["A", 2]}, {"field": "", "op": ""}]},
        ]}
        f = self.j("post", "/api/filters", {"name": "Alt", "state": {"q": "", "adv": adv}}, status=201)
        self.assertEqual(f["state"]["adv"], {"mode": "any", "items": [
            {"field": "@reviewed_at", "op": "older", "value": "365"},
            {"mode": "none", "items": [{"field": "keywords", "op": "any", "value": ["A", "2"]}, {"field": "", "op": ""}]},
        ]})
        self.assertNotIn("adv", self.j("post", "/api/filters", {"name": "Ohne", "state": {"q": "x"}}, status=201)["state"])
        deep = {"mode": "all", "items": []}
        for _ in range(5):
            deep = {"mode": "all", "items": [deep]}
        for bad in [{"mode": "xor", "items": []}, {"mode": "all", "items": "x"}, {"mode": "all", "items": [1]},
                    {"mode": "all", "items": [{"field": "a", "op": "eq"}] * 51}, deep, "nope"]:
            self.j("post", "/api/filters", {"name": "x", "state": {"adv": bad}}, status=400)


class RecentChangesTests(ApiTestBase):
    def test_maintainers_only_and_hidden_fields(self):
        f = FieldDef.objects.get(key="contract_no")
        f.visible_to = "admin"
        f.save()
        self.as_("admin")
        r = self.j("post", "/api/records", {"kind": "person", "last_name": "Meier",
                                            "data": {"contract_no": "geheim", "email": "m@x.de"}}, status=201)
        AuditEntry.objects.create(username="admin", action="update", target_type="user", target_id="1")
        self.as_("read")
        self.j("get", "/api/recent-changes", status=403)
        self.as_("maintainer")
        entries = self.j("get", "/api/recent-changes")["entries"]
        self.assertEqual(len(entries), 1)  # record entries only
        e = entries[0]
        self.assertEqual((e["target_id"], e["kind"], e["action"]), (r["uid"], "person", "create"))
        self.assertIn("email", e["changes"])
        self.assertNotIn("contract_no", e["changes"])


class ExportListTests(ApiTestBase):
    def setUp(self):
        super().setUp()
        f = FieldDef.objects.get(key="contract_no")
        f.visible_to = "maintainer"
        f.save()
        self.users["admin"].first_name, self.users["admin"].last_name = "Ada", "Admin"
        self.users["admin"].save()
        self.c = Record.objects.create(kind="company", name="ACME", data={"keywords": ["A", "B"]})
        self.p = Record.objects.create(kind="person", first_name="Max", last_name="Muster", company=self.c,
                                       data={"email": "m@acme.de", "contract_no": "geheim",
                                             "internal_contact": "admin"})
        self.gone = Record.objects.create(kind="person", last_name="Weg", deleted_at="2026-01-01T00:00:00Z")

    def test_csv_respects_order_and_visibility(self):
        self.as_("read")
        resp = self.client.post("/api/export", {"format": "csv", "uids": [str(self.p.uid), str(self.c.uid),
                                                                          str(self.gone.uid), "junk"]},
                                content_type="application/json")
        self.assertEqual(resp.status_code, 200)
        self.assertIn('filename="kartei-', resp["Content-Disposition"])
        rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig")), delimiter=";"))
        header = rows[0]
        self.assertEqual(header[:4], ["Art", "Nachname", "Vorname", "Firma"])
        self.assertNotIn("Vertragsnummer", header)
        self.assertEqual(len(rows), 3)
        person = dict(zip(header, rows[1]))
        self.assertEqual((person["Art"], person["Nachname"], person["Firma"]), ("Person", "Muster", "ACME"))
        self.assertEqual(person["E-Mail"], "m@acme.de")
        self.assertEqual(person["Ansprechpartner intern"], "Ada Admin")
        company = dict(zip(header, rows[2]))
        self.assertEqual((company["Art"], company["Firma"], company["Schlagworte"]), ("Firma", "ACME", "A, B"))
        e = AuditEntry.objects.get(action="export")
        self.assertEqual((e.username, e.target_type, e.changes["rows"]), ("read", "export", [None, 2]))

    def test_xlsx_english(self):
        Profile.objects.filter(user=self.users["maintainer"]).update(language="en")
        self.as_("maintainer")
        resp = self.client.post("/api/export", {"format": "xlsx", "uids": [str(self.p.uid)]},
                                content_type="application/json")
        ws = load_workbook(io.BytesIO(resp.content)).active
        header = [c.value for c in ws[1]]
        self.assertEqual(header[:4], ["Type", "Last name", "First name", "Company"])
        self.assertIn("Contract no.", header)
        self.assertEqual(ws.max_row, 2)

    def test_invalid(self):
        self.as_("read")
        self.j("post", "/api/export", {"format": "pdf", "uids": [str(self.p.uid)]}, status=400)
        self.j("post", "/api/export", {"format": "csv", "uids": []}, status=400)


class BulkTests(ApiTestBase):
    def setUp(self):
        super().setUp()
        FieldDef.objects.filter(key="keywords").update(options=[{"value": v, "de": v, "en": v} for v in "AB"])
        self.c = Record.objects.create(kind="company", name="ACME", data={"keywords": ["A"]})
        self.p1 = Record.objects.create(kind="person", last_name="Eins", data={"keywords": ["A", "B"]})
        self.p2 = Record.objects.create(kind="person", last_name="Zwei")
        self.uids = [str(r.uid) for r in (self.c, self.p1, self.p2)]

    def bulk(self, status=200, **body):
        return self.j("post", "/api/records/bulk", {"uids": self.uids, **body}, status=status)

    def test_add_remove_keywords(self):
        self.as_("read")
        self.bulk(op="add", key="keywords", value=["C"], status=403)
        self.as_("maintainer")
        out = self.bulk(op="add", key="keywords", value=["C", "a"])
        self.assertEqual({r["uid"]: r["data"].get("keywords") for r in out["records"]},
                         {self.uids[0]: ["A", "C"], self.uids[1]: ["A", "B", "C"], self.uids[2]: ["C", "A"]})  # "a" maps onto option "A"
        self.assertEqual(out["errors"], [])
        self.assertIn("C", [o["value"] for o in FieldDef.objects.get(key="keywords").options])  # open options
        self.p1.refresh_from_db()
        self.assertEqual(self.p1.version, 2)
        e = AuditEntry.objects.filter(target_id=self.uids[1], action="update").first()
        self.assertEqual(e.changes, {"keywords": [["A", "B"], ["A", "B", "C"]], "bulk": [None, True]})
        out = self.bulk(op="remove", key="keywords", value=["A", "B"])
        self.assertEqual([r["data"]["keywords"] for r in out["records"]], [["C"], ["C"], ["C"]])

    def test_set_clear_and_errors(self):
        self.as_("maintainer")
        out = self.bulk(op="set", key="job_title", value="Chef")
        self.assertEqual([r["uid"] for r in out["records"]], self.uids[1:])
        self.assertEqual(out["errors"], [{"uid": self.uids[0], "name": "ACME", "fields": {"job_title": "not_applicable"}}])
        bad = self.bulk(op="set", key="email", value="kaputt")
        self.assertEqual(bad["records"], [])
        self.assertEqual({e["fields"]["email"] for e in bad["errors"]}, {"invalid_email", "not_applicable"})
        out = self.bulk(op="set", key="internal_contact", value="read")  # the responsible user
        self.assertEqual({r["data"]["internal_contact"] for r in out["records"]}, {"read"})
        self.bulk(op="clear", key="internal_contact")
        self.assertNotIn("internal_contact", Record.objects.get(pk=self.p1.pk).data)
        FieldDef.objects.filter(key="internal_contact").update(required=True)
        out = self.bulk(op="clear", key="job_title")
        self.assertEqual({e["fields"].get("internal_contact") for e in out["errors"] if e["name"] != "ACME"}, {"required"})

    def test_rejected_requests(self):
        self.as_("maintainer")
        FieldDef.objects.filter(key="contract_no").update(editable_by="admin")
        self.assertEqual(self.bulk(op="set", key="contract_no", value="x", status=400)["fields"], {"key": "forbidden"})
        self.assertEqual(self.bulk(op="add", key="job_title", value=["x"], status=400)["fields"], {"key": "invalid"})
        self.assertEqual(self.bulk(op="set", key="nope", value="x", status=400)["fields"], {"key": "unknown_field"})
        self.assertEqual(self.bulk(op="set", key="job_title", value="", status=400)["fields"], {"value": "required"})
        self.bulk(op="explode", key="job_title", status=400)

    def test_delete(self):
        self.as_("maintainer")
        out = self.bulk(op="delete")
        self.assertEqual(out["deleted"], self.uids)
        self.assertEqual(Record.objects.filter(deleted_at__isnull=True).count(), 0)
        self.assertEqual(AuditEntry.objects.filter(action="delete", changes__bulk=[None, True]).count(), 3)
        self.assertEqual(self.bulk(op="delete")["errors"][0]["fields"], {"record": "not_found"})
